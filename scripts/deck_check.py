#!/usr/bin/env python3
"""coulis deck check: run the gate over a deck, one slide at a time, then look across slides.

    python3 scripts/deck_check.py deck.pdf            # or .pptx, .md (slides split on '---' or '# ')
    python3 scripts/deck_check.py deck.pdf --profile ~/.coulis

Per slide: the usual verdict with each issue quoted. Across slides: the same thing said with two
different numbers (e.g. "300 customers" on slide 3, "350 customers" on slide 9), which is the
mistake investors catch first. No model call.
"""
import argparse
import json
import os
import re
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import gate  # noqa: E402
from coulis_core.profile import profile_from_folder  # noqa: E402

NOUN = re.compile(r"\b(\$?\d[\d,.]*\s?(?:k|m|b)?\+?%?)\s*(customers|users|founders|companies|teams|posts|followers|arr|mrr|revenue|"
                  r"installs|downloads|members|readers|subscribers|employees|countries|markets|pilots|tests|calls|meetings)\b", re.I)


def slides_from(path):
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        import pymupdf
        d = pymupdf.open(p)
        out = [pg.get_text() for pg in d]
        d.close()
        return out
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for s in Presentation(p).slides:
            out.append("\n".join(sh.text for sh in s.shapes if hasattr(sh, "text")))
        return out
    text = p.read_text()
    parts = re.split(r"^---\s*$|^(?=# )", text, flags=re.M)
    return [t for t in parts if t.strip()]


def cross_slide(slides):
    """Same noun, different numbers on different slides."""
    seen = defaultdict(dict)
    for i, s in enumerate(slides, 1):
        for num, noun in NOUN.findall(s):
            seen[noun.lower()].setdefault(num.replace(",", "").lower(), i)
    out = []
    for noun, nums in seen.items():
        if len(nums) > 1:
            out.append({"noun": noun, "values": {n: f"slide {i}" for n, i in nums.items()}})
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--profile", default=os.path.expanduser("~/.coulis"))
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    prof = Path(a.profile)
    if not (prof / "claims.md").exists():
        prof = Path(os.path.expanduser("~/.tastee"))
    persona = profile_from_folder(prof).to_persona()
    slides = slides_from(a.deck)
    report = {"deck": a.deck, "slides": len(slides), "per_slide": [], "cross_slide": cross_slide(slides)}
    for i, s in enumerate(slides, 1):
        if not s.strip():
            continue
        r = gate.check(s, persona, "linkedin")
        report["per_slide"].append({"slide": i, "verdict": r["verdict"], "sameness": r["sameness"]["score"],
                                    "issues": [{"level": x["level"], "check": x["check"], "quote": x["quote"]} for x in r["issues"]]})
    worst = "BLOCK" if any(x["verdict"] == "BLOCK" for x in report["per_slide"]) else "FIX" if any(x["verdict"] == "FIX" for x in report["per_slide"]) or report["cross_slide"] else "PASS"
    report["verdict"] = worst
    if a.json:
        print(json.dumps(report, indent=2))
        return
    print(f"{worst}  {len(slides)} slides")
    for x in report["per_slide"]:
        if x["issues"]:
            print(f"  slide {x['slide']}: {x['verdict']}")
            for i in x["issues"]:
                print(f"     {i['level']} {i['check']}: {i['quote']!r}")
    for c in report["cross_slide"]:
        print(f"  across slides: '{c['noun']}' appears as " + ", ".join(f"{n} ({w})" for n, w in c["values"].items()))


if __name__ == "__main__":
    main()
